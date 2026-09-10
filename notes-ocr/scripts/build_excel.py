"""
notes-ocr / build_excel.py
Reusable Excel builder for digitized handwritten notes.

Usage (standalone):
    python build_excel.py --data structured_data.json --out output.xlsx

Or import and call build_workbook(data, output_path) directly from another script.

Expected data structure (dict):
{
  "title": "Document Title",
  "author": "Name or null",
  "date": "DD-MM-YYYY or null",
  "sections": [
    {
      "name": "Section Name",       # becomes sheet tab name (max 31 chars)
      "description": "...",         # optional prose summary
      "fields": [                   # key-value pairs → two-column sheet
        {"label": "Field", "value": "Value"},
        ...
      ],
      "tables": [                   # structured tables
        {
          "title": "Table Name",
          "headers": ["Col1", "Col2", ...],
          "rows": [["val", "val", ...], ...]
        }
      ],
      "diagrams": [                 # diagram references
        {"name": "Diagram Name", "description": "What it shows"}
      ]
    }
  ]
}
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl not found — run: pip install openpyxl --break-system-packages")


# ── Style constants ───────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
HEADER_FONT   = Font(bold=True, color="FFFFFF", size=11, name="Arial")
SUBHDR_FILL   = PatternFill(start_color="B4C7E7", end_color="B4C7E7", fill_type="solid")
SUBHDR_FONT   = Font(bold=True, size=10, name="Arial")
BODY_FONT     = Font(size=10, name="Arial")
ALT_FILL      = PatternFill(start_color="F0F4FA", end_color="F0F4FA", fill_type="solid")
THIN_BORDER   = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin")
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="top",    wrap_text=True)


def _apply_header(cell, value):
    cell.value = value
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = CENTER
    cell.border = THIN_BORDER


def _apply_body(cell, value, alt=False):
    cell.value = value
    cell.font = BODY_FONT
    cell.fill = ALT_FILL if alt else PatternFill()
    cell.alignment = LEFT
    cell.border = THIN_BORDER


def _auto_width(ws, max_width=60):
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 4
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width, max_width)


def _safe_sheet_name(name, used):
    """Truncate to 31 chars and deduplicate."""
    name = str(name)[:31]
    original = name
    i = 2
    while name in used:
        suffix = f" ({i})"
        name = original[:31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def build_workbook(data: dict, output_path: str):
    wb = Workbook()
    wb.remove(wb.active)
    used_names = set()

    # ── Sheet 1: Overview ────────────────────────────────────────────────────
    ws_ov = wb.create_sheet(_safe_sheet_name("Overview", used_names))
    ws_ov.column_dimensions["A"].width = 22
    ws_ov.column_dimensions["B"].width = 55

    _apply_header(ws_ov.cell(1, 1), "Field")
    _apply_header(ws_ov.cell(1, 2), "Value")

    overview_rows = [
        ("Document Title", data.get("title", "")),
        ("Author",         data.get("author", "—")),
        ("Date",           data.get("date", "—")),
        ("Sections",       str(len(data.get("sections", [])))),
    ]
    for r, (lbl, val) in enumerate(overview_rows, 2):
        _apply_body(ws_ov.cell(r, 1), lbl, alt=(r % 2 == 0))
        _apply_body(ws_ov.cell(r, 2), val, alt=(r % 2 == 0))

    # Section list
    ws_ov.cell(len(overview_rows) + 3, 1).value = "Section"
    ws_ov.cell(len(overview_rows) + 3, 1).font = SUBHDR_FONT
    ws_ov.cell(len(overview_rows) + 3, 1).fill = SUBHDR_FILL
    ws_ov.cell(len(overview_rows) + 3, 2).value = "Description"
    ws_ov.cell(len(overview_rows) + 3, 2).font = SUBHDR_FONT
    ws_ov.cell(len(overview_rows) + 3, 2).fill = SUBHDR_FILL

    for i, sec in enumerate(data.get("sections", []), len(overview_rows) + 4):
        _apply_body(ws_ov.cell(i, 1), sec.get("name", ""), alt=(i % 2 == 0))
        _apply_body(ws_ov.cell(i, 2), sec.get("description", ""), alt=(i % 2 == 0))

    # ── One sheet per section ────────────────────────────────────────────────
    for sec in data.get("sections", []):
        ws = wb.create_sheet(_safe_sheet_name(sec.get("name", "Section"), used_names))
        row = 1

        # Section title banner
        ws.merge_cells(f"A{row}:D{row}")
        title_cell = ws.cell(row, 1, sec.get("name", ""))
        title_cell.font = Font(bold=True, color="FFFFFF", size=12, name="Arial")
        title_cell.fill = HEADER_FILL
        title_cell.alignment = CENTER
        row += 1

        # Description prose
        if sec.get("description"):
            ws.merge_cells(f"A{row}:D{row}")
            desc_cell = ws.cell(row, 1, sec["description"])
            desc_cell.font = BODY_FONT
            desc_cell.alignment = LEFT
            ws.row_dimensions[row].height = max(30, len(sec["description"]) // 3)
            row += 2

        # Key-value fields
        fields = sec.get("fields", [])
        if fields:
            ws.cell(row, 1).value = "Field"
            ws.cell(row, 1).font = SUBHDR_FONT
            ws.cell(row, 1).fill = SUBHDR_FILL
            ws.cell(row, 2).value = "Value"
            ws.cell(row, 2).font = SUBHDR_FONT
            ws.cell(row, 2).fill = SUBHDR_FILL
            row += 1
            for i, f in enumerate(fields):
                _apply_body(ws.cell(row, 1), f.get("label", ""), alt=(i % 2 == 0))
                _apply_body(ws.cell(row, 2), f.get("value", ""), alt=(i % 2 == 0))
                row += 1
            row += 1

        # Tables
        for tbl in sec.get("tables", []):
            # Table title
            headers = tbl.get("headers", [])
            n_cols = max(len(headers), 1)
            end_col = get_column_letter(n_cols)
            ws.merge_cells(f"A{row}:{end_col}{row}")
            tbl_title = ws.cell(row, 1, tbl.get("title", "Table"))
            tbl_title.font = Font(bold=True, color="FFFFFF", size=10, name="Arial")
            tbl_title.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            tbl_title.alignment = CENTER
            row += 1

            # Headers
            for c, h in enumerate(headers, 1):
                _apply_header(ws.cell(row, c), h)
            row += 1

            # Rows
            for i, data_row in enumerate(tbl.get("rows", [])):
                for c, val in enumerate(data_row, 1):
                    _apply_body(ws.cell(row, c), val, alt=(i % 2 == 0))
                row += 1
            row += 2

        # Diagrams references
        for diag in sec.get("diagrams", []):
            ws.merge_cells(f"A{row}:D{row}")
            ws.cell(row, 1).value = f"📊 Diagram: {diag.get('name', '')}"
            ws.cell(row, 1).font = Font(bold=True, size=10, name="Arial", color="366092")
            row += 1
            ws.merge_cells(f"A{row}:D{row}")
            ws.cell(row, 1).value = diag.get("description", "")
            ws.cell(row, 1).font = BODY_FONT
            ws.cell(row, 1).alignment = LEFT
            ws.merge_cells(f"A{row+1}:D{row+1}")
            ws.cell(row+1, 1).value = "→ See Mermaid diagram in the Markdown (.md) file"
            ws.cell(row+1, 1).font = Font(italic=True, size=9, name="Arial", color="595959")
            row += 3

        _auto_width(ws)

    wb.save(output_path)
    print(f"✅ Excel saved: {output_path}")
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Build Excel from structured notes data")
    parser.add_argument("--data", required=True, help="Path to structured_data.json")
    parser.add_argument("--out",  required=True, help="Output .xlsx path")
    args = parser.parse_args()

    with open(args.data) as f:
        data = json.load(f)

    build_workbook(data, args.out)


if __name__ == "__main__":
    main()
